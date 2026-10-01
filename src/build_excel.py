"""
Step 9 - Build the Excel version of the model (excel/Supply_Chain_Analytics.xlsx).

Unlike reports/supply_chain_report.xlsx (a values-only export), this workbook
is a LIVE model: every result is an Excel formula. Change an assumption
(service level, ordering cost, risk weights, alpha...) and everything updates.

Run:  python src/build_excel.py      (after python src/main.py has created data/)

Excel version vs Python version
  * Safety stock uses the std-dev of weekly demand (last 52 weeks) instead of
    the forecast error, because that is the standard spreadsheet approach.
    Results are therefore close to, but not identical with, the Python numbers.
  * Forecast sheet back-tests one SKU at a time with one-step-ahead forecasts.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import config
from inventory import SEGMENT_STRATEGY

OUT = config.ROOT / "excel" / "Supply_Chain_Analytics.xlsx"

# ---------- styles ----------
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="123D73")
F_SUB = Font(name=FONT, size=10, italic=True, color="52514E")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_KPI = Font(name=FONT, size=18, bold=True, color="123D73")
FILL_HEAD = PatternFill("solid", fgColor="1D5AA6")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_KPI = PatternFill("solid", fgColor="EEF5FD")
FILL_SECTION = PatternFill("solid", fgColor="F0EFEA")
THIN = Side(style="thin", color="E1E0D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {
    "Critical - expedite": "F6C9C9", "Order now": "FBDCCD",
    "Healthy": "D3F0D3", "Overstock": "FEEDC4",
}

FMT_MONEY = '"CAD "#,##0;("CAD "#,##0);"-"'
FMT_MONEY2 = '"CAD "#,##0.00;("CAD "#,##0.00);"-"'
FMT_INT = '#,##0;(#,##0);"-"'
FMT_1 = '#,##0.0;(#,##0.0);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def style_range(ws, rng, font=F_BASE, fmt=None):
    for row in ws[rng]:
        for c in row:
            c.font = font
            if fmt:
                c.number_format = fmt


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


# ======================================================================
def build_landed_cost(wb, ws, products, suppliers, s_last):
    """Formula-driven offshore vs nearshore comparison (mirrors src/landed_cost.py)."""
    title(ws, "Logistics & Landed Cost: Offshore vs Nearshore",
          "Total annual landed cost = units x landed unit cost + carrying cost of safety stock and in-transit inventory. "
          "Blue-on-yellow cells are inputs.")
    hdr = ["Country", "Price index", "Freight %", "Duty %", "Brokerage %", "Region", "Landed multiplier",
           "Cost to land 100", "Avg lead time (wks)"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(4, i, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True)
    r0 = 5
    for k, (country, (pi, fr, du, br, region)) in enumerate(config.LANDED_COST.items()):
        r = r0 + k
        ws.cell(r, 1, country).font = F_BASE
        for col, val in ((2, pi), (3, fr), (4, du), (5, br)):
            c = ws.cell(r, col, val)
            c.font, c.fill = F_INPUT, FILL_INPUT
            c.number_format = "0.00" if col == 2 else FMT_PCT
        ws.cell(r, 6, region).font = F_BASE
        ws.cell(r, 7, f"=B{r}*(1+C{r}+D{r}+E{r})").number_format = "0.000"
        ws.cell(r, 8, f"=100*G{r}").number_format = "0.0"
        ws.cell(r, 9, f'=IFERROR(AVERAGEIF(Supplier_Scorecard!$C$5:$C${s_last},A{r},Supplier_Scorecard!$I$5:$I${s_last}),"-")').number_format = "0.0"
    r1 = r0 + len(config.LANDED_COST) - 1
    name(wb, "LC_Country", f"Landed_Cost!$A${r0}:$A${r1}")
    name(wb, "LC_Mult", f"Landed_Cost!$G${r0}:$G${r1}")
    name(wb, "LC_Price", f"Landed_Cost!$B${r0}:$B${r1}")
    ws.cell(r1 + 1, 1, "Duty rates are illustrative (CUSMA, CETA, CPTPP). Check the Canada Customs Tariff for real HS codes.").font = F_SUB
    ws.cell(r1 + 2, 1, "Accept nearshore premium up to").font = F_BOLD
    c = ws.cell(r1 + 2, 2, config.NEARSHORE_PREMIUM_OK)
    c.font, c.fill, c.number_format = F_INPUT, FILL_INPUT, FMT_PCT
    name(wb, "LC_Premium", f"Landed_Cost!$B${r1 + 2}")

    # nearshore supplier per country (first supplier found in that country)
    near = {}
    for country in ("Mexico", "USA", "Canada"):
        m = suppliers[suppliers["country"] == country]
        if len(m):
            near[country] = m.iloc[0]["supplier_id"]
    country_of = dict(zip(suppliers["supplier_id"], suppliers["country"]))
    offshore = products[products["supplier_id"].map(lambda s: country_of.get(s) not in (None, "Mexico", "USA", "Canada")
                                                    and country_of.get(s) in config.LANDED_COST)]

    H = r1 + 5
    cols = ["SKU", "Product", "Supplier", "Country", "Weekly demand", "Demand std", "Unit cost now", "Base price",
            "Now: landed unit", "Now: inventory value", "Now: total/yr"]
    for c_name in near:
        cols += [f"{c_name}: total/yr", f"{c_name}: inventory value"]
    cols += ["Best nearshore", "Best nearshore total/yr", "Saving/yr", "Saving %", "Inventory released", "Recommendation", "row"]
    for i, h in enumerate(cols, start=1):
        c = ws.cell(H, i, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[H].height = 42
    ws.cell(H - 1, 1, "Comparison for every SKU bought offshore").font = F_BOLD
    idx = {h: i + 1 for i, h in enumerate(cols)}
    L = get_column_letter

    def lt(sup_ref):
        return f"INDEX(SC_LT,MATCH({sup_ref},SC_ID,0))", f"INDEX(SC_LTStd,MATCH({sup_ref},SC_ID,0))"

    for k, p in enumerate(offshore.itertuples()):
        r = H + 1 + k
        R = f"$" + L(idx["row"]) + str(r)
        ws.cell(r, idx["row"], f'=MATCH("{p.sku}",P_SKU,0)').font = Font(name=FONT, size=8, color="A0A0A0")
        g = lambda rng: f"=INDEX({rng},{R})"
        ws.cell(r, idx["SKU"], g("P_SKU"))
        ws.cell(r, idx["Product"], g("P_Name"))
        ws.cell(r, idx["Supplier"], g("P_Sup"))
        sup = f"{L(idx['Supplier'])}{r}"
        ws.cell(r, idx["Country"], f"=INDEX(Supplier_Scorecard!$C$5:$C${s_last},MATCH({sup},SC_ID,0))")
        ws.cell(r, idx["Weekly demand"], g("P_Demand"))
        ws.cell(r, idx["Demand std"], g("P_Std"))
        ws.cell(r, idx["Unit cost now"], g("P_Cost"))
        d, sd = f"{L(idx['Weekly demand'])}{r}", f"{L(idx['Demand std'])}{r}"
        cty = f"{L(idx['Country'])}{r}"
        ws.cell(r, idx["Base price"], f"={L(idx['Unit cost now'])}{r}/INDEX(LC_Price,MATCH({cty},LC_Country,0))")
        base = f"{L(idx['Base price'])}{r}"

        def option(country_ref, sup_ref, unit_cell=None):
            unit = unit_cell or f"({base}*INDEX(LC_Mult,MATCH({country_ref},LC_Country,0)))"
            ltm, lts = lt(sup_ref)
            inv = f"(z*SQRT({ltm}*{sd}^2+{d}^2*{lts}^2)+{d}*{ltm})*{unit}"
            total = f"{d}*WeeksYear*{unit}+{inv}*HoldRate"
            return inv, total

        ws.cell(r, idx["Now: landed unit"], f"={base}*INDEX(LC_Mult,MATCH({cty},LC_Country,0))")
        unit_now = f"{L(idx['Now: landed unit'])}{r}"
        inv, total = option(cty, sup, unit_now)
        ws.cell(r, idx["Now: inventory value"], f"={inv}")
        ws.cell(r, idx["Now: total/yr"], f"={total}")
        tot_cells, inv_cells, names = [], [], []
        for c_name, sid in near.items():
            inv, total = option(f'"{c_name}"', f'"{sid}"')
            tc = ws.cell(r, idx[f"{c_name}: total/yr"], f"=IFERROR({total},9.99E+99)")
            ic = ws.cell(r, idx[f"{c_name}: inventory value"], f"=IFERROR({inv},0)")
            tot_cells.append(f"{L(tc.column)}{r}")
            inv_cells.append(f"{L(ic.column)}{r}")
            names.append(c_name)
        best_tot = "MIN(" + ",".join(tot_cells) + ")"
        choose_name = "CHOOSE(MATCH(" + best_tot + ",(" + ",".join(tot_cells) + "),0)," if False else None
        # pick name / inventory of the cheapest option with nested IFs (works in every Excel version)
        def pick(values):
            expr = values[-1]
            for tc_, v in zip(reversed(tot_cells[:-1]), reversed(values[:-1])):
                expr = f"IF({tc_}={best_tot},{v},{expr})"
            return expr
        ws.cell(r, idx["Best nearshore"], "=" + pick([f'"{n}"' for n in names]))
        ws.cell(r, idx["Best nearshore total/yr"], f"={best_tot}")
        now_t, best_t = f"{L(idx['Now: total/yr'])}{r}", f"{L(idx['Best nearshore total/yr'])}{r}"
        ws.cell(r, idx["Saving/yr"], f"={now_t}-{best_t}")
        sav = f"{L(idx['Saving/yr'])}{r}"
        ws.cell(r, idx["Saving %"], f"=IFERROR({sav}/{now_t},0)")
        pct = f"{L(idx['Saving %'])}{r}"
        ws.cell(r, idx["Inventory released"], f"={L(idx['Now: inventory value'])}{r}-" + pick(inv_cells))
        ws.cell(r, idx["Recommendation"],
                f'=IF({sav}>0,"Nearshore: lower total cost",IF({pct}>=-LC_Premium,"Nearshore candidate: small premium for resilience","Keep offshore"))')
        for col in range(1, idx["row"]):
            ws.cell(r, col).font = F_LINK if col <= 3 else F_BASE
            ws.cell(r, col).border = BOX
        for h in cols:
            if h in ("Weekly demand", "Demand std"):
                ws.cell(r, idx[h]).number_format = FMT_1
            elif h in ("Unit cost now", "Base price", "Now: landed unit"):
                ws.cell(r, idx[h]).number_format = FMT_MONEY2
            elif "total" in h or "inventory" in h.lower() or h in ("Saving/yr",):
                ws.cell(r, idx[h]).number_format = FMT_MONEY
            elif h == "Saving %":
                ws.cell(r, idx[h]).number_format = FMT_PCT
    last = H + len(offshore)
    rec_col = L(idx["Recommendation"])
    for txt, color in (("Nearshore: lower total cost", "D3F0D3"), ("Nearshore candidate", "FEEDC4"), ("Keep offshore", "F0EFEA")):
        ws.conditional_formatting.add(f"{rec_col}{H + 1}:{rec_col}{last}",
                                      FormulaRule(formula=[f'LEFT({rec_col}{H + 1},{len(txt)})="{txt}"'],
                                                  fill=PatternFill("solid", fgColor=color)))
    # summary
    S = r0  # summary sits to the right of the country table, above the SKU table
    sv, rc, iv = L(idx["Saving/yr"]), rec_col, L(idx["Inventory released"])
    rng = lambda col: f"{col}{H + 1}:{col}{last}"
    summary = [("Offshore SKUs reviewed", f"=COUNTA({rng(L(idx['SKU']))})", FMT_INT),
               ("Cheaper to source nearshore", f'=COUNTIF({rng(rc)},"Nearshore: lower total cost")', FMT_INT),
               ("Nearshore candidates (incl. small premium)", f'=COUNTIF({rng(rc)},"Nearshore*")', FMT_INT),
               ("Saving/yr if cheaper SKUs move", f'=SUMIF({rng(sv)},">0")', FMT_MONEY),
               ("Inventory released if candidates move", f'=SUMIF({rng(rc)},"Nearshore*",{rng(iv)})', FMT_MONEY)]
    ws.cell(S - 1, 11, "Summary").font = F_BOLD
    for k, (lab, f_, fmt) in enumerate(summary):
        ws.cell(S + k, 11, lab).font = F_BASE
        c = ws.cell(S + k, 14, f_)
        c.font, c.number_format, c.fill = F_BOLD, fmt, FILL_KPI
    widths = {1: 12, 2: 26, 3: 10, 4: 11}
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[L(i)].width = widths.get(i, 13)
    ws.column_dimensions[rec_col].width = 44
    ws.freeze_panes = ws.cell(H + 1, 3)


def build(source_note: str = ""):
    d = config.DATA_DIR
    products = pd.read_csv(d / "products.csv")
    suppliers = pd.read_csv(d / "suppliers.csv")
    demand = pd.read_csv(d / "weekly_demand.csv", parse_dates=["week_start"]).sort_values(["sku", "week_start"])
    pos = pd.read_csv(d / "purchase_orders.csv", parse_dates=["order_date", "promised_date", "received_date"])
    stock = pd.read_csv(d / "inventory_snapshot.csv")
    suppliers = suppliers[suppliers["supplier_id"].isin(pos["supplier_id"])].reset_index(drop=True)
    weeks = sorted(demand["week_start"].unique())

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_dash = wb.create_sheet("Dashboard")
    ws_a = wb.create_sheet("Assumptions")
    ws_pol = wb.create_sheet("Inventory_Policy")
    ws_sc = wb.create_sheet("Supplier_Scorecard")
    ws_fc = wb.create_sheet("Forecast")
    ws_sim = wb.create_sheet("Disruption_Simulator")
    ws_lc = wb.create_sheet("Landed_Cost")
    ws_p = wb.create_sheet("Products")
    ws_s = wb.create_sheet("Suppliers")
    ws_d = wb.create_sheet("Weekly_Demand")
    ws_po = wb.create_sheet("Purchase_Orders")
    ws_inv = wb.create_sheet("Inventory")

    nP, nS, nD, nPO = len(products), len(suppliers), len(demand), len(pos)
    pEnd, sEnd, dEnd, poEnd = nP + 1, nS + 1, nD + 1, nPO + 1

    # ------------------------------------------------------------------ data sheets
    header(ws_p, 1, ["SKU", "Product", "Category", "Unit cost (CAD)", "Supplier ID"], [12, 34, 18, 14, 12])
    for r, p in enumerate(products.itertuples(), start=2):
        ws_p.append([p.sku, p.product_name, p.category, p.unit_cost_cad, p.supplier_id])
    style_range(ws_p, f"A2:E{pEnd}")
    style_range(ws_p, f"D2:D{pEnd}", fmt=FMT_MONEY2)

    header(ws_s, 1, ["Supplier ID", "Supplier", "Country", "Promised lead time (wks)"], [12, 32, 14, 16])
    for s in suppliers.itertuples():
        ws_s.append([s.supplier_id, s.supplier_name, s.country, s.promised_lead_time_wks])
    style_range(ws_s, f"A2:D{sEnd}")

    header(ws_d, 1, ["Week start", "SKU", "Units", "In last 52 wks (1/0)", "Units² (last 52 wks)"], [12, 12, 10, 14, 16])
    for r, row in enumerate(demand.itertuples(), start=2):
        ws_d.cell(r, 1, row.week_start.to_pydatetime()).number_format = "yyyy-mm-dd"
        ws_d.cell(r, 2, row.sku)
        ws_d.cell(r, 3, int(row.units))
        ws_d.cell(r, 4, f"=IF(A{r}>LastWeek-364,1,0)")
        ws_d.cell(r, 5, f"=C{r}^2*D{r}")
    style_range(ws_d, f"A2:E{dEnd}")
    for c in ws_d[f"A2:A{dEnd}"]:
        c[0].number_format = "yyyy-mm-dd"

    header(ws_po, 1, ["PO ID", "Supplier ID", "SKU", "Order date", "Promised date", "Received date",
                      "Qty ordered", "Qty received", "Qty defective", "PO value (CAD)",
                      "Lead time (wks)", "Lead time²", "On time (1/0)", "Days late"],
           [11, 11, 11, 12, 12, 12, 11, 11, 11, 13, 11, 11, 11, 10])
    for r, p in enumerate(pos.itertuples(), start=2):
        ws_po.cell(r, 1, p.po_id); ws_po.cell(r, 2, p.supplier_id); ws_po.cell(r, 3, p.sku)
        for col, val in ((4, p.order_date), (5, p.promised_date), (6, p.received_date)):
            ws_po.cell(r, col, val.to_pydatetime()).number_format = "yyyy-mm-dd"
        ws_po.cell(r, 7, int(p.qty_ordered)); ws_po.cell(r, 8, int(p.qty_received))
        ws_po.cell(r, 9, int(p.qty_defective)); ws_po.cell(r, 10, float(p.po_value_cad))
        ws_po.cell(r, 11, f"=(F{r}-D{r})/7")
        ws_po.cell(r, 12, f"=K{r}^2")
        ws_po.cell(r, 13, f"=IF(F{r}<=E{r},1,0)")
        ws_po.cell(r, 14, f"=MAX(0,F{r}-E{r})")
    style_range(ws_po, f"A2:N{poEnd}")
    for col in "DEF":
        for c in ws_po[f"{col}2:{col}{poEnd}"]:
            c[0].number_format = "yyyy-mm-dd"
    style_range(ws_po, f"J2:J{poEnd}", fmt=FMT_MONEY2)
    style_range(ws_po, f"K2:L{poEnd}", fmt="0.00")

    header(ws_inv, 1, ["SKU", "On hand (units)", "On order (units)"], [12, 14, 14])
    for s in stock.itertuples():
        ws_inv.append([s.sku, int(s.on_hand_units), int(s.on_order_units)])
    style_range(ws_inv, f"A2:C{len(stock) + 1}")

    for ws in (ws_p, ws_s, ws_d, ws_po, ws_inv):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # ------------------------------------------------------------------ assumptions
    title(ws_a, "Assumptions", "Blue text on yellow = inputs you can change. Everything else in the workbook recalculates.")
    ws_a.column_dimensions["A"].width = 38
    ws_a.column_dimensions["B"].width = 16
    ws_a.column_dimensions["C"].width = 70
    rows = [
        ("Inventory policy", None, None, None),
        ("Service level", config.SERVICE_LEVEL, "ServiceLevel", "Target share of order cycles with no stockout"),
        ("z-score", "=NORMSINV(ServiceLevel)", "z", "Calculated from the service level (95% -> 1.645)"),
        ("Ordering cost per PO (CAD)", config.ORDER_COST, "OrderCost", "Admin, freight booking, receiving. Assumption"),
        ("Annual holding cost rate", config.HOLDING_COST_RATE, "HoldRate", "Capital, storage, insurance, obsolescence. Assumption"),
        ("Weeks per year", config.WEEKS_PER_YEAR, "WeeksYear", ""),
        ("Current ordering habit (weeks of demand per PO)", 4, "CurrentCoverWks", "Used to compare today's ordering with EOQ"),
        ("Last week in data", f"=MAX(Weekly_Demand!$A$2:$A${dEnd})", "LastWeek", "Analysis uses the 52 weeks up to this date"),
        ("Classification", None, None, None),
        ("ABC: A-class cut-off (cumulative spend)", config.ABC_A_CUTOFF, "ACut", "First 80% of spend = A"),
        ("ABC: B-class cut-off", config.ABC_B_CUTOFF, "BCut", "Next 15% = B, rest = C"),
        ("XYZ: X max coefficient of variation", config.XYZ_X_MAX_CV, "XMaxCV", "Stable demand"),
        ("XYZ: Y max coefficient of variation", config.XYZ_Y_MAX_CV, "YMaxCV", "Variable demand; above = Z (erratic)"),
        ("Supplier risk weights (must total 100%)", None, None, None),
        ("Weight: late delivery", config.RISK_WEIGHTS["late_delivery"], "wLate", ""),
        ("Weight: lead-time variability", config.RISK_WEIGHTS["lead_time_variability"], "wLTV", ""),
        ("Weight: quality (defects)", config.RISK_WEIGHTS["quality"], "wQuality", ""),
        ("Weight: country / route risk", config.RISK_WEIGHTS["country_risk"], "wCountry", ""),
        ("Weight: spend dependency", config.RISK_WEIGHTS["spend_dependency"], "wSpend", ""),
        ("Check: weights total", "=wLate+wLTV+wQuality+wCountry+wSpend", None, "Should be 100%"),
        ("'Clearly bad' thresholds (factor score = 1)", None, None, None),
        ("Late-delivery rate", 0.35, "BadLate", "35% of POs late"),
        ("Lead-time coefficient of variation", 0.30, "BadLTV", ""),
        ("Defect rate", 0.025, "BadQuality", "2.5% of units defective"),
        ("Share of total spend with one supplier", 0.30, "BadSpend", ""),
        ("Disruption simulator", None, None, None),
        ("Selling price markup on cost", 1.35, "Markup", "Used to turn units short into revenue at risk. Assumption"),
    ]
    r = 4
    for label, val, nm, note in rows:
        if val is None and nm is None:
            ws_a.cell(r, 1, label).font = F_BOLD
            for col in range(1, 4):
                ws_a.cell(r, col).fill = FILL_SECTION
        else:
            ws_a.cell(r, 1, label).font = F_BASE
            c = ws_a.cell(r, 2, val)
            is_input = not (isinstance(val, str) and val.startswith("="))
            c.font = F_INPUT if is_input else F_BASE
            if is_input:
                c.fill = FILL_INPUT
            if nm:
                name(wb, nm, f"Assumptions!$B${r}")
            ws_a.cell(r, 3, note).font = F_SUB
            if isinstance(val, float) and val < 1 and nm not in ("z",) or nm in ("ServiceLevel",) or label.startswith("Check"):
                c.number_format = FMT_PCT
            if nm == "z":
                c.number_format = "0.000"
            if nm == "LastWeek":
                c.number_format = "yyyy-mm-dd"
            if nm == "OrderCost":
                c.number_format = FMT_MONEY2
            if nm == "Markup":
                c.number_format = "0.00x"
        r += 1

    # country risk table
    r += 1
    ws_a.cell(r, 1, "Country / route risk (1 = low, 5 = high)").font = F_BOLD
    for col in range(1, 4):
        ws_a.cell(r, col).fill = FILL_SECTION
    r += 1
    c_start = r
    for country, score in config.COUNTRY_RISK.items():
        ws_a.cell(r, 1, country).font = F_BASE
        c = ws_a.cell(r, 2, score)
        c.font, c.fill = F_INPUT, FILL_INPUT
        r += 1
    ws_a.cell(c_start, 3, "Illustrative assumptions (distance, port congestion, tariffs, geopolitics), not official ratings").font = F_SUB
    name(wb, "CountryList", f"Assumptions!$A${c_start}:$A${r - 1}")
    name(wb, "CountryScore", f"Assumptions!$B${c_start}:$B${r - 1}")

    # action table
    r += 1
    ws_a.cell(r, 1, "Recommended action by main risk driver").font = F_BOLD
    for col in range(1, 4):
        ws_a.cell(r, col).fill = FILL_SECTION
    r += 1
    a_start = r
    actions = [
        ("late delivery", "Quarterly business review with OTD target; add late-delivery clause to contract"),
        ("lead time variability", "Raise safety stock on its SKUs; book freight earlier; track via ASN"),
        ("quality", "Supplier corrective action request (SCAR); increase incoming inspection"),
        ("country risk", "Qualify a nearshore/domestic second source for critical SKUs"),
        ("spend dependency", "Dual-source top SKUs to reduce concentration risk"),
    ]
    for drv, act in actions:
        ws_a.cell(r, 1, drv).font = F_BASE
        ws_a.cell(r, 3, act).font = F_BASE
        r += 1
    name(wb, "DriverList", f"Assumptions!$A${a_start}:$A${r - 1}")
    name(wb, "DriverAction", f"Assumptions!$C${a_start}:$C${r - 1}")

    # segment strategy table
    r += 1
    ws_a.cell(r, 1, "Stocking strategy by ABC-XYZ segment").font = F_BOLD
    for col in range(1, 4):
        ws_a.cell(r, col).fill = FILL_SECTION
    r += 1
    g_start = r
    for seg, strat in SEGMENT_STRATEGY.items():
        ws_a.cell(r, 1, seg).font = F_BASE
        ws_a.cell(r, 3, strat).font = F_BASE
        r += 1
    name(wb, "SegList", f"Assumptions!$A${g_start}:$A${r - 1}")
    name(wb, "SegStrategy", f"Assumptions!$C${g_start}:$C${r - 1}")

    # handy whole-column names for data
    name(wb, "D_Week", f"Weekly_Demand!$A$2:$A${dEnd}")
    name(wb, "D_SKU", f"Weekly_Demand!$B$2:$B${dEnd}")
    name(wb, "D_Units", f"Weekly_Demand!$C$2:$C${dEnd}")
    name(wb, "D_Last52", f"Weekly_Demand!$D$2:$D${dEnd}")
    name(wb, "D_Units2", f"Weekly_Demand!$E$2:$E${dEnd}")
    for col, nm in zip("BCGHIJKLMN", ["PO_Sup", "PO_SKU", "PO_QtyOrd", "PO_QtyRec", "PO_QtyDef",
                                      "PO_Value", "PO_LT", "PO_LT2", "PO_OnTime", "PO_DaysLate"]):
        name(wb, nm, f"Purchase_Orders!${col}$2:${col}${poEnd}")

    # ------------------------------------------------------------------ supplier scorecard
    sc = ws_sc
    title(sc, "Supplier Scorecard & Risk Score", "Every column is a formula over the Purchase_Orders sheet. Weights and thresholds are on the Assumptions sheet.")
    H = 4
    sc_cols = ["Supplier ID", "Supplier", "Country", "POs", "Spend (CAD)", "Share of spend",
               "On-time delivery", "Avg days late", "Avg lead time (wks)", "Lead time std (wks)",
               "Lead time CV", "Fill rate", "Defect rate", "Country risk (1-5)",
               "Factor: late", "Factor: LT variability", "Factor: quality", "Factor: country", "Factor: spend",
               "late delivery", "lead time variability", "quality", "country risk", "spend dependency",
               "Risk score (0-100)", "Risk tier", "Main risk driver", "Recommended action", "Risk rank"]
    header(sc, H, sc_cols, [11, 30, 11, 7, 14, 10, 10, 9, 10, 10, 9, 9, 9, 9,
                            9, 10, 9, 9, 9, 9, 10, 9, 9, 10, 11, 9, 20, 60, 8])
    sc.cell(H - 1, 20, "Weighted contribution of each factor (used to find the main driver)").font = F_SUB
    s0, s1 = H + 1, H + nS
    for i in range(nS):
        r = H + 1 + i
        sr = i + 2
        f = {
            1: f"=Suppliers!A{sr}", 2: f"=Suppliers!B{sr}", 3: f"=Suppliers!C{sr}",
            4: f"=COUNTIF(PO_Sup,A{r})",
            5: f"=SUMIF(PO_Sup,A{r},PO_Value)",
            6: f"=IFERROR(E{r}/SUM($E${s0}:$E${s1}),0)",
            7: f"=IFERROR(AVERAGEIF(PO_Sup,A{r},PO_OnTime),0)",
            8: f"=IFERROR(AVERAGEIF(PO_Sup,A{r},PO_DaysLate),0)",
            9: f"=IFERROR(AVERAGEIF(PO_Sup,A{r},PO_LT),0)",
            10: f"=IFERROR(SQRT(MAX(0,(SUMIF(PO_Sup,A{r},PO_LT2)-D{r}*I{r}^2)/(D{r}-1))),0)",
            11: f"=IFERROR(J{r}/I{r},0)",
            12: f"=IFERROR(SUMIF(PO_Sup,A{r},PO_QtyRec)/SUMIF(PO_Sup,A{r},PO_QtyOrd),0)",
            13: f"=IFERROR(SUMIF(PO_Sup,A{r},PO_QtyDef)/SUMIF(PO_Sup,A{r},PO_QtyRec),0)",
            14: f"=IFERROR(INDEX(CountryScore,MATCH(C{r},CountryList,0)),3)",
            15: f"=MIN(1,MAX(0,(1-G{r})/BadLate))",
            16: f"=MIN(1,MAX(0,K{r}/BadLTV))",
            17: f"=MIN(1,MAX(0,M{r}/BadQuality))",
            18: f"=MIN(1,MAX(0,(N{r}-1)/4))",
            19: f"=MIN(1,MAX(0,F{r}/BadSpend))",
            20: f"=O{r}*wLate", 21: f"=P{r}*wLTV", 22: f"=Q{r}*wQuality",
            23: f"=R{r}*wCountry", 24: f"=S{r}*wSpend",
            25: f"=ROUND(SUM(T{r}:X{r})*100,1)",
            26: f'=IF(Y{r}>60,"High",IF(Y{r}>35,"Medium","Low"))',
            27: f"=INDEX($T${H}:$X${H},MATCH(MAX(T{r}:X{r}),T{r}:X{r},0))",
            28: (f'=IF(AND(Z{r}="Low",F{r}>=0.25),"Reliable, but holds a large share of spend: keep as primary and qualify a backup supplier for its A-class SKUs",'
                 f'IF(Z{r}="Low","Maintain; candidate for volume consolidation / preferred-supplier status",'
                 f'IF(Z{r}="High","Priority: ","")&INDEX(DriverAction,MATCH(AA{r},DriverList,0))))'),
            29: f"=RANK(Y{r},$Y${s0}:$Y${s1})",
        }
        for col, formula in f.items():
            c = sc.cell(r, col, formula)
            c.font = F_LINK if col <= 3 else F_BASE
            c.border = BOX
        for col in (6, 7, 12, 13, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 11):
            sc.cell(r, col).number_format = FMT_PCT if col in (6, 7, 11, 12) else ("0.00%" if col == 13 else "0.00")
        sc.cell(r, 5).number_format = FMT_MONEY
        for col in (8, 9, 10):
            sc.cell(r, col).number_format = "0.0"
        sc.cell(r, 25).number_format = "0.0"
        sc.cell(r, 28).alignment = Alignment(wrap_text=True, vertical="top")
    sc.freeze_panes = sc.cell(H + 1, 3)
    tier_rng = f"Z{s0}:Z{s1}"
    for tier, color in (("High", "F6C9C9"), ("Medium", "FEEDC4"), ("Low", "D3F0D3")):
        sc.conditional_formatting.add(tier_rng, CellIsRule(operator="equal", formula=[f'"{tier}"'],
                                                           fill=PatternFill("solid", fgColor=color)))
    name(wb, "SC_ID", f"Supplier_Scorecard!$A${s0}:$A${s1}")
    name(wb, "SC_LT", f"Supplier_Scorecard!$I${s0}:$I${s1}")
    name(wb, "SC_LTStd", f"Supplier_Scorecard!$J${s0}:$J${s1}")

    # ------------------------------------------------------------------ inventory policy
    pol = ws_pol
    title(pol, "Inventory Policy: ABC-XYZ, Safety Stock, Reorder Point, EOQ",
          "One row per SKU. Demand statistics use the last 52 weeks. Status compares inventory position (on hand + on order) with safety stock and reorder point.")
    H = 4
    pol_cols = ["SKU", "Product", "Category", "Supplier ID", "Unit cost (CAD)",
                "Annual units (52w)", "Avg weekly demand", "Weekly demand std", "Annual spend (CAD)",
                "Cum. spend share before SKU", "ABC", "Demand CV", "XYZ", "Segment",
                "Lead time (wks)", "Lead time std (wks)", "Safety stock", "Reorder point", "EOQ",
                "On hand", "On order", "Inventory position", "Weeks of cover", "Status",
                "Suggested order qty", "Excess value (CAD)", "Safety stock value (CAD)",
                "Current ordering cost/yr", "EOQ ordering cost/yr", "Annual saving (CAD)", "Strategy",
                "helper: simulator row", "helper: action row"]
    header(pol, H, pol_cols, [11, 30, 16, 9, 10, 10, 10, 10, 13, 10, 6, 8, 6, 8, 9, 9, 9, 9, 9,
                              9, 9, 10, 9, 17, 10, 12, 12, 12, 12, 12, 55, 9, 9])
    p0, p1 = H + 1, H + nP
    for i in range(nP):
        r = H + 1 + i
        pr = i + 2
        f = {
            1: f"=Products!A{pr}", 2: f"=Products!B{pr}", 3: f"=Products!C{pr}", 4: f"=Products!E{pr}",
            5: f"=Products!D{pr}",
            6: f"=SUMIFS(D_Units,D_SKU,A{r},D_Last52,1)",
            7: f"=IFERROR(AVERAGEIFS(D_Units,D_SKU,A{r},D_Last52,1),0)",
            8: (f"=IFERROR(SQRT(MAX(0,(SUMIFS(D_Units2,D_SKU,A{r})-COUNTIFS(D_SKU,A{r},D_Last52,1)*G{r}^2)"
                f"/(COUNTIFS(D_SKU,A{r},D_Last52,1)-1))),0)"),
            9: f"=F{r}*E{r}",
            10: f'=IFERROR(SUMIF($I${p0}:$I${p1},">"&I{r})/SUM($I${p0}:$I${p1}),0)',
            11: f'=IF(J{r}<ACut,"A",IF(J{r}<BCut,"B","C"))',
            12: f"=IFERROR(H{r}/G{r},0)",
            13: f'=IF(L{r}<=XMaxCV,"X",IF(L{r}<=YMaxCV,"Y","Z"))',
            14: f"=K{r}&M{r}",
            15: f"=IFERROR(INDEX(SC_LT,MATCH(D{r},SC_ID,0)),0)",
            16: f"=IFERROR(INDEX(SC_LTStd,MATCH(D{r},SC_ID,0)),0)",
            17: f"=ROUNDUP(z*SQRT(O{r}*H{r}^2+G{r}^2*P{r}^2),0)",
            18: f"=ROUNDUP(G{r}*O{r}+Q{r},0)",
            19: f"=IFERROR(ROUNDUP(SQRT(2*G{r}*WeeksYear*OrderCost/(E{r}*HoldRate)),0),0)",
            20: f"=IFERROR(INDEX(Inventory!$B:$B,MATCH(A{r},Inventory!$A:$A,0)),0)",
            21: f"=IFERROR(INDEX(Inventory!$C:$C,MATCH(A{r},Inventory!$A:$A,0)),0)",
            22: f"=T{r}+U{r}",
            23: f"=IFERROR(T{r}/G{r},0)",
            24: f'=IF(V{r}<Q{r},"Critical - expedite",IF(V{r}<=R{r},"Order now",IF(V{r}>R{r}+S{r},"Overstock","Healthy")))',
            25: f'=IF(OR(X{r}="Critical - expedite",X{r}="Order now"),MAX(S{r},R{r}+S{r}-V{r}),0)',
            26: f'=IF(X{r}="Overstock",(V{r}-R{r}-S{r})*E{r},0)',
            27: f"=Q{r}*E{r}",
            28: (f"=IFERROR(G{r}*WeeksYear/(G{r}*CurrentCoverWks)*OrderCost"
                 f"+G{r}*CurrentCoverWks/2*E{r}*HoldRate,0)"),
            29: f"=IFERROR(G{r}*WeeksYear/S{r}*OrderCost+S{r}/2*E{r}*HoldRate,0)",
            30: f"=AB{r}-AC{r}",
            31: f"=IFERROR(INDEX(SegStrategy,MATCH(N{r},SegList,0)),\"\")",
            32: f'=IF(D{r}=SimSupplier,COUNTIF($D${p0}:D{r},SimSupplier),"")',
            33: f'=IF(OR(X{r}="Critical - expedite",X{r}="Order now"),COUNTIF($X${p0}:X{r},"Critical - expedite")+COUNTIF($X${p0}:X{r},"Order now"),"")',
        }
        for col, formula in f.items():
            c = pol.cell(r, col, formula)
            c.font = F_LINK if col <= 5 else F_BASE
            c.border = BOX
        pol.cell(r, 5).number_format = FMT_MONEY2
        for col in (6, 17, 18, 19, 20, 21, 22, 25):
            pol.cell(r, col).number_format = FMT_INT
        for col in (7, 8, 15, 16, 23):
            pol.cell(r, col).number_format = FMT_1
        for col in (9, 26, 27, 28, 29, 30):
            pol.cell(r, col).number_format = FMT_MONEY
        pol.cell(r, 10).number_format = FMT_PCT
        pol.cell(r, 12).number_format = "0.00"
    pol.freeze_panes = pol.cell(H + 1, 3)
    pol.auto_filter.ref = f"A{H}:AG{p1}"
    for status, color in STATUS_FILL.items():
        pol.conditional_formatting.add(f"X{p0}:X{p1}", CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                                  fill=PatternFill("solid", fgColor=color)))
    for nm, col in (("P_SKU", "A"), ("P_Name", "B"), ("P_Sup", "D"), ("P_Cost", "E"), ("P_Spend", "I"),
                    ("P_ABC", "K"), ("P_XYZ", "M"), ("P_Seg", "N"), ("P_LT", "O"), ("P_SS", "Q"),
                    ("P_ROP", "R"), ("P_EOQ", "S"), ("P_OnHand", "T"), ("P_OnOrder", "U"), ("P_Pos", "V"),
                    ("P_Demand", "G"), ("P_Std", "H"), ("P_Status", "X"), ("P_Order", "Y"), ("P_Excess", "Z"),
                    ("P_SSValue", "AA"), ("P_CurCost", "AB"), ("P_Saving", "AD"),
                    ("P_SimRow", "AF"), ("P_ActRow", "AG")):
        name(wb, nm, f"Inventory_Policy!${col}${p0}:${col}${p1}")

    # ------------------------------------------------------------------ forecast sheet
    fc = ws_fc
    title(fc, "Demand Forecast Back-test (one SKU at a time)",
          "Pick a SKU. Three methods forecast each week one step ahead; accuracy (WAPE) is measured on the last 12 weeks.")
    labels = [("SKU (pick from list)", products["sku"].iloc[0], "FcSKU"),
              ("Product", "=IFERROR(INDEX(P_Name,MATCH(FcSKU,P_SKU,0)),\"\")", None),
              ("Smoothing factor alpha (0-1)", 0.3, "Alpha"),
              ("Test window (weeks)", 12, None)]
    for i, (lab, val, nm) in enumerate(labels):
        r = 4 + i
        fc.cell(r, 1, lab).font = F_BOLD
        c = fc.cell(r, 2, val)
        if not (isinstance(val, str) and val.startswith("=")):
            c.font, c.fill = F_INPUT, FILL_INPUT
        else:
            c.font = F_BASE
        if nm:
            name(wb, nm, f"Forecast!$B${r}")
    dv = DataValidation(type="list", formula1=f"=Products!$A$2:$A${pEnd}", allow_blank=False)
    fc.add_data_validation(dv)
    dv.add("B4")
    fc.column_dimensions["A"].width = 26
    fc.column_dimensions["B"].width = 14
    for col in "CDEFGH":
        fc.column_dimensions[col].width = 14

    T0 = 16
    hist_n = len(weeks)
    first, last = T0 + 1, T0 + hist_n
    test_first = last - 11
    fut_first, fut_last = last + 1, last + 12
    header_labels = ["Week #", "Week start", "Actual units", "Moving avg (4w)", "SES level",
                     "Exp. smoothing", "Seasonal naive", "Best method"]
    for i, lab in enumerate(header_labels, start=1):
        c = fc.cell(T0, i, lab)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True)

    # accuracy block
    fc.cell(9, 1, "Accuracy on last 12 weeks").font = F_BOLD
    acc = [("Moving avg (4w)", "D"), ("Exp. smoothing", "F"), ("Seasonal naive", "G")]
    for i, (lab, col) in enumerate(acc):
        r = 10 + i
        fc.cell(r, 1, lab).font = F_BASE
        c = fc.cell(r, 2, f"=IFERROR(SUMPRODUCT(ABS($C${test_first}:$C${last}-{col}{test_first}:{col}{last}))/SUM($C${test_first}:$C${last}),0)")
        c.number_format, c.font = FMT_PCT, F_BASE
    fc.cell(13, 1, "Best method (lowest WAPE)").font = F_BOLD
    fc["B13"] = "=INDEX(A10:A12,MATCH(MIN(B10:B12),B10:B12,0))"
    fc["B13"].font = F_BOLD
    fc.cell(14, 1, "Next 12 weeks, best method (units)").font = F_BOLD
    fc["B14"] = f"=SUM(H{fut_first}:H{fut_last})"
    fc["B14"].number_format = FMT_INT
    fc["C10"] = "WAPE = sum(|actual - forecast|) / sum(actual)"
    fc["C10"].font = F_SUB

    for k in range(hist_n):
        r = first + k
        fc.cell(r, 1, k + 1)
        fc.cell(r, 2, pd.Timestamp(weeks[k]).to_pydatetime()).number_format = "yyyy-mm-dd"
        fc.cell(r, 3, f"=SUMIFS(D_Units,D_SKU,FcSKU,D_Week,B{r})")
        fc.cell(r, 4, f"=AVERAGE(C{r - 4}:C{r - 1})" if k >= 4 else None)
        fc.cell(r, 5, f"=C{r}" if k == 0 else f"=Alpha*C{r}+(1-Alpha)*E{r - 1}")
        fc.cell(r, 6, f"=E{r - 1}" if k >= 1 else None)
        fc.cell(r, 7, f"=C{r - 52}" if k >= 52 else None)
        fc.cell(r, 8, f'=IF(k_best=1,D{r},IF(k_best=2,F{r},G{r}))' if k >= hist_n - 12 else None)
    for j in range(12):
        r = fut_first + j
        fc.cell(r, 1, hist_n + j + 1)
        fc.cell(r, 2, f"=B{r - 1}+7").number_format = "yyyy-mm-dd"
        fc.cell(r, 4, f"=AVERAGE($C${last - 3}:$C${last})")
        fc.cell(r, 6, f"=$E${last}")
        fc.cell(r, 7, f"=C{r - 52}")
        fc.cell(r, 8, f"=IF(k_best=1,D{r},IF(k_best=2,F{r},G{r}))")
    name(wb, "k_best", "Forecast!$B$15")
    fc["A15"] = "Best method index (1-3)"
    fc["A15"].font = F_SUB
    fc["B15"] = "=MATCH(MIN(B10:B12),B10:B12,0)"
    style_range(fc, f"A{first}:H{fut_last}")
    for r in range(first, fut_last + 1):
        fc.cell(r, 2).number_format = "yyyy-mm-dd"
        for col in range(3, 9):
            fc.cell(r, col).number_format = "#,##0"
        if r >= fut_first:
            for col in range(1, 9):
                fc.cell(r, col).fill = FILL_KPI
    fc.cell(fut_first - 1, 9, "<- history ends; shaded rows = 12-week forecast").font = F_SUB

    # Scatter chart with lines: each series has its own date range, so the
    # forecast line starts where the test window starts (no blank/zero gaps)
    ch = ScatterChart()
    ch.title = "Actual demand vs best-method forecast"
    ch.style = 13
    ch.y_axis.title = "Units per week"
    ch.height, ch.width = 8, 22
    s_act = Series(Reference(fc, min_col=3, min_row=first, max_row=last),
                   Reference(fc, min_col=2, min_row=first, max_row=last), title="Actual units")
    s_fc = Series(Reference(fc, min_col=8, min_row=test_first, max_row=fut_last),
                  Reference(fc, min_col=2, min_row=test_first, max_row=fut_last), title="Forecast (best method)")
    for s, color, dash in ((s_act, "2A78D6", None), (s_fc, "EB6834", "dash")):
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 20000
        if dash:
            s.graphicalProperties.line.dashStyle = dash
        ch.series.append(s)
    ch.x_axis.number_format = "mmm-yy"
    ch.x_axis.majorUnit = 91
    ch.x_axis.title = "Week"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    fc.add_chart(ch, "D3")
    fc.freeze_panes = fc.cell(T0 + 1, 1)

    # ------------------------------------------------------------------ disruption simulator
    sim = ws_sim
    title(sim, "Disruption Simulator", "What if a supplier is delayed (port strike, tariff hold, factory shutdown)? Pick a supplier and a delay.")
    sim["A4"], sim["A5"] = "Supplier (pick from list)", "Extra delay (weeks)"
    sim["B4"], sim["B5"] = suppliers["supplier_id"].iloc[-1], 4
    for c in ("B4", "B5"):
        sim[c].font, sim[c].fill = F_INPUT, FILL_INPUT
    sim["A4"].font = sim["A5"].font = F_BOLD
    sim["C4"] = "=IFERROR(INDEX(Supplier_Scorecard!$B$5:$B$" + str(s1) + ",MATCH(B4,SC_ID,0)),\"\")"
    name(wb, "SimSupplier", "Disruption_Simulator!$B$4")
    name(wb, "SimDelay", "Disruption_Simulator!$B$5")
    dv2 = DataValidation(type="list", formula1=f"=Suppliers!$A$2:$A${sEnd}", allow_blank=False)
    sim.add_data_validation(dv2)
    dv2.add("B4")
    sim["A7"], sim["A8"], sim["A9"] = "SKUs that stock out", "Revenue at risk (CAD)", "Extra buffer stock to prevent it (CAD)"
    for c in ("A7", "A8", "A9"):
        sim[c].font = F_BOLD
    SH = 12
    max_rows = max(30, int(products["supplier_id"].value_counts().max()) + 2)
    s_first, s_last = SH + 1, SH + max_rows
    sim["B7"] = f'=COUNTIF(J{s_first}:J{s_last},"Stockout")&" of "&COUNTIF(P_Sup,SimSupplier)'
    sim["B8"] = f"=SUM(L{s_first}:L{s_last})"
    sim["B9"] = f"=SUM(N{s_first}:N{s_last})"
    sim["B8"].number_format = sim["B9"].number_format = FMT_MONEY
    for c in ("B7", "B8", "B9"):
        sim[c].font, sim[c].fill = F_KPI, FILL_KPI
    header_sim = ["#", "SKU", "Product", "ABC", "Avg weekly demand", "On hand + on order",
                  "Weeks of stock", "Lead time (wks)", "New lead time", "Outcome", "Units short",
                  "Revenue at risk (CAD)", "Buffer units", "Buffer cost (CAD)", "row"]
    for i, lab in enumerate(header_sim, start=1):
        c = sim.cell(SH, i, lab)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True)
    widths = [5, 11, 30, 6, 11, 12, 10, 10, 10, 11, 10, 14, 10, 13, 6]
    for i, w in enumerate(widths, start=1):
        sim.column_dimensions[get_column_letter(i)].width = w
    sim.column_dimensions["A"].width = 34
    for k in range(max_rows):
        r = SH + 1 + k
        sim.cell(r, 1, k + 1)
        sim.cell(r, 15, f'=IFERROR(MATCH(A{r},P_SimRow,0),"")')
        g = lambda rng: f'=IF(O{r}="","",INDEX({rng},O{r}))'
        sim.cell(r, 2, g("P_SKU")); sim.cell(r, 3, g("P_Name")); sim.cell(r, 4, g("P_ABC"))
        sim.cell(r, 5, g("P_Demand"))
        sim.cell(r, 6, f'=IF(O{r}="","",INDEX(P_Pos,O{r}))')
        sim.cell(r, 7, f'=IF(O{r}="","",IFERROR(F{r}/E{r},999))')
        sim.cell(r, 8, g("P_LT"))
        sim.cell(r, 9, f'=IF(O{r}="","",H{r}+SimDelay)')
        sim.cell(r, 10, f'=IF(O{r}="","",IF(I{r}>G{r},"Stockout","Covered"))')
        sim.cell(r, 11, f'=IF(O{r}="","",ROUNDUP(MAX(0,I{r}-G{r})*E{r},0))')
        sim.cell(r, 12, f'=IF(O{r}="","",K{r}*INDEX(P_Cost,O{r})*Markup)')
        sim.cell(r, 13, f'=IF(O{r}="","",ROUNDUP(E{r}*SimDelay,0))')
        sim.cell(r, 14, f'=IF(O{r}="","",M{r}*INDEX(P_Cost,O{r}))')
        for col in range(1, 16):
            sim.cell(r, col).font = F_BASE
        for col in (5, 7, 8, 9):
            sim.cell(r, col).number_format = FMT_1
        for col in (6, 11, 13):
            sim.cell(r, col).number_format = FMT_INT
        for col in (12, 14):
            sim.cell(r, col).number_format = FMT_MONEY
        sim.cell(r, 15).font = Font(name=FONT, size=8, color="A0A0A0")
    sim.conditional_formatting.add(f"J{s_first}:J{s_last}", CellIsRule(operator="equal", formula=['"Stockout"'], fill=PatternFill("solid", fgColor="F6C9C9")))
    sim.conditional_formatting.add(f"J{s_first}:J{s_last}", CellIsRule(operator="equal", formula=['"Covered"'], fill=PatternFill("solid", fgColor="D3F0D3")))

    # ------------------------------------------------------------------ landed cost
    build_landed_cost(wb, ws_lc, products, suppliers, s1)

    # ------------------------------------------------------------------ dashboard
    db = ws_dash
    db.sheet_view.showGridLines = False
    title(db, "Supply Chain Control Tower",
          "Inventory optimization, demand forecasting and supplier risk. All figures are live formulas; change inputs on the Assumptions sheet.")
    for col, w in zip("ABCDEFGHIJKLMNOP", [3, 22, 14, 3, 22, 14, 3, 22, 14, 3, 22, 14, 3, 14, 14, 14]):
        db.column_dimensions[col].width = w
    kpis = [
        ("B4", "Annual spend analysed", "=SUM(P_Spend)", FMT_MONEY, '=COUNTA(P_SKU)&" SKUs, "&COUNTA(SC_ID)&" suppliers"'),
        ("E4", "Spend in A-class SKUs", '=SUMIF(P_ABC,"A",P_Spend)/SUM(P_Spend)', FMT_PCT, '="from "&COUNTIF(P_ABC,"A")&" SKUs ("&TEXT(COUNTIF(P_ABC,"A")/COUNTA(P_SKU),"0%")&")"'),
        ("H4", "SKUs to order now", '=COUNTIF(P_Status,"Order now")+COUNTIF(P_Status,"Critical - expedite")', FMT_INT, '=COUNTIF(P_Status,"Critical - expedite")&" critical (below safety stock)"'),
        ("K4", "EOQ saving per year", "=SUM(P_Saving)", FMT_MONEY, '=TEXT(SUM(P_Saving)/SUM(P_CurCost),"0.0%")&" lower ordering + holding cost"'),
        ("B8", "Excess inventory", "=SUM(P_Excess)", FMT_MONEY, '=COUNTIF(P_Status,"Overstock")&" SKUs above max stock"'),
        ("E8", "Safety stock investment", "=SUM(P_SSValue)", FMT_MONEY, '="at "&TEXT(ServiceLevel,"0%")&" service level"'),
        ("H8", "Supplier on-time delivery", "=AVERAGE(PO_OnTime)", FMT_PCT, '=COUNTA(PO_OnTime)&" purchase orders"'),
        ("K8", "High-risk suppliers", '=COUNTIF(Supplier_Scorecard!$Z$5:$Z$' + str(s1) + ',"High")', FMT_INT, '="Top: "&INDEX(Supplier_Scorecard!$B$5:$B$' + str(s1) + ',MATCH(1,Supplier_Scorecard!$AC$5:$AC$' + str(s1) + ',0))'),
    ]
    for cell, label, formula, fmt, ctx in kpis:
        col = db[cell].column
        row = db[cell].row
        db.cell(row, col, label).font = Font(name=FONT, size=9, color="52514E")
        v = db.cell(row + 1, col, formula)
        v.font, v.number_format = F_KPI, fmt
        v.alignment = Alignment(horizontal="left")
        db.cell(row + 2, col, ctx).font = Font(name=FONT, size=8, color="7A7974")
        for rr in range(row, row + 3):
            db.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + 1)
        db.row_dimensions[row + 1].height = 26
        for rr in range(row, row + 3):
            for cc in (col, col + 1):
                db.cell(rr, cc).fill = FILL_KPI

    # status table
    db["B12"] = "Inventory status"
    db["B12"].font = F_BOLD
    for i, st in enumerate(["Critical - expedite", "Order now", "Healthy", "Overstock"]):
        r = 13 + i
        db.cell(r, 2, st).font = F_BASE
        db.cell(r, 2).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
        db.cell(r, 3, f'=COUNTIF(P_Status,B{r})').font = F_BASE
    bc = BarChart()
    bc.type = "bar"
    bc.title = "SKUs by inventory status"
    bc.legend = None
    bc.add_data(Reference(db, min_col=3, min_row=13, max_row=16), titles_from_data=False)
    bc.set_categories(Reference(db, min_col=2, min_row=13, max_row=16))
    bc.series[0].graphicalProperties.solidFill = "EC835A"
    bc.height, bc.width = 6.5, 7
    bc.y_axis.majorGridlines = None
    db.add_chart(bc, "B18")

    # ABC-XYZ matrix
    hdrs = ["X stable", "Y variable", "Z erratic"]
    db["M12"] = "ABC-XYZ matrix (SKU count)"
    db["M12"].font = F_BOLD
    for j, h in enumerate(hdrs):
        c = db.cell(13, 14 + j, h)
        c.font, c.fill = F_HEAD, FILL_HEAD
    for i, a in enumerate("ABC"):
        r = 14 + i
        c = db.cell(r, 13, a)
        c.font, c.fill = F_HEAD, FILL_HEAD
        for j, x in enumerate("XYZ"):
            cell = db.cell(r, 14 + j, f'=COUNTIFS(P_ABC,"{a}",P_XYZ,"{x}")')
            cell.font, cell.border = F_BOLD, BOX
            cell.alignment = Alignment(horizontal="center")
    from openpyxl.formatting.rule import ColorScaleRule
    db.conditional_formatting.add("N14:P16", ColorScaleRule(start_type="min", start_color="EEF5FD", end_type="max", end_color="4F92DE"))

    # supplier risk chart
    sr = Reference(sc, min_col=25, min_row=4, max_row=s1)
    rc = BarChart()
    rc.type = "bar"
    rc.title = "Supplier risk score (0-100)"
    rc.legend = None
    rc.add_data(sr, titles_from_data=True)
    rc.set_categories(Reference(sc, min_col=2, min_row=5, max_row=s1))
    rc.series[0].graphicalProperties.solidFill = "D03B3B"
    rc.y_axis.scaling.min, rc.y_axis.scaling.max = 0, 100
    rc.y_axis.majorGridlines = None
    rc.height, rc.width = 9, 15
    db.add_chart(rc, "E12")

    # action list
    AL = 33
    db.cell(AL, 2, "What to order today").font = F_BOLD
    act = [("Product", 2, "P_Name"), ("SKU", 3, "P_SKU"), ("Status", 5, "P_Status"), ("Position", 6, "P_Pos"),
           ("Reorder point", 8, "P_ROP"), ("Order qty", 9, "P_Order"), ("Supplier", 11, "P_Sup")]
    for lab, col, _ in act:
        c = db.cell(AL + 1, col, lab)
        c.font, c.fill = F_HEAD, FILL_HEAD
    n_act = 20
    HC = 18  # hidden helper column R
    for k in range(n_act):
        r = AL + 2 + k
        db.cell(r, HC, f'=IFERROR(MATCH({k + 1},P_ActRow,0),"")')
        for lab, col, rng in act:
            c = db.cell(r, col, f'=IF($R{r}="","",INDEX({rng},$R{r}))')
            c.font = F_BASE
            if lab in ("Position", "Reorder point", "Order qty"):
                c.number_format = FMT_INT
    db.column_dimensions["R"].hidden = True
    for status, color in STATUS_FILL.items():
        db.conditional_formatting.add(f"E{AL + 2}:E{AL + 1 + n_act}", CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)))
    db.cell(AL + 2 + n_act, 2, '=IF(COUNT(P_ActRow)>' + str(n_act) + ',"+ "&(COUNT(P_ActRow)-' + str(n_act) + ')&" more on the Inventory_Policy sheet","")').font = F_SUB

    # ------------------------------------------------------------------ README sheet
    rd = ws_readme
    rd.sheet_view.showGridLines = False
    rd.column_dimensions["A"].width = 26
    rd.column_dimensions["B"].width = 100
    title(rd, "Supply Chain Inventory Optimization & Supplier Risk Analytics (Excel model)",
          "Every result is a live Excel formula. Built by Dhruvilkumar Upadhyay.")
    lines = [
        ("How to use", ""),
        ("1. Assumptions", "Change any blue-on-yellow cell (service level, ordering cost, risk weights...). The whole model recalculates."),
        ("2. Dashboard", "KPIs, inventory status, supplier risk chart, ABC-XYZ matrix and the 'what to order today' list."),
        ("3. Inventory_Policy", "One row per SKU: ABC-XYZ class, safety stock, reorder point, EOQ, status and suggested order quantity."),
        ("4. Supplier_Scorecard", "On-time delivery, lead time, defects, spend share and a weighted 0-100 risk score with a recommended action."),
        ("5. Forecast", "Pick a SKU from the dropdown in B4 to back-test 3 forecasting methods and see a 12-week forecast."),
        ("6. Disruption_Simulator", "Pick a supplier (B4) and an extra delay (B5) to see which SKUs stock out and the revenue at risk."),
        ("7. Landed_Cost", "Offshore vs nearshore: total landed cost (price, freight, duty, brokerage, plus the cost of safety stock and in-transit inventory) for every SKU bought overseas."),
        ("Data sheets", "Products, Suppliers, Weekly_Demand, Purchase_Orders, Inventory. Replace them with your own data (same columns)."),
        ("Colour code", "Blue text = input you can change. Green text = link to another sheet. Black = formula."),
        ("Key formulas", ""),
        ("Safety stock", "z x SQRT( lead time x demand std² + avg demand² x lead time std² )"),
        ("Reorder point", "avg weekly demand x lead time + safety stock"),
        ("EOQ", "SQRT( 2 x annual demand x ordering cost / (unit cost x holding rate) )"),
        ("WAPE", "SUM(|actual - forecast|) / SUM(actual)"),
        ("Data source", source_note),
        ("Python version", "The same analysis is automated in Python (src/main.py). That version uses forecast error for safety stock, so numbers differ slightly."),
    ]
    for i, (a, b) in enumerate(lines):
        r = 4 + i
        ca, cb = rd.cell(r, 1, a), rd.cell(r, 2, b)
        ca.font = F_BOLD if b == "" or a[0].isdigit() or a in ("Data sheets", "Colour code", "Safety stock", "Reorder point", "EOQ", "WAPE", "Data source", "Python version") else F_BASE
        cb.font = F_BASE
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if b == "":
            ca.fill = cb.fill = FILL_SECTION

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"README": "52514E", "Dashboard": "1D5AA6", "Assumptions": "FAB219"}.get(ws.title, None)
    for ws in (ws_dash, ws_readme, ws_sim):   # print each of these on one landscape page
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
    wb.calculation.fullCalcOnLoad = True   # Excel computes every formula when the file opens
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT.relative_to(config.ROOT)}")


if __name__ == "__main__":
    raw = config.KAGGLE_DATACO_FILE
    use_kaggle = config.DATA_SOURCE == "kaggle" or (config.DATA_SOURCE == "auto" and raw.exists())
    if use_kaggle:
        import load_kaggle_dataco as src
    else:
        import generate_data as src
    build(src.SOURCE_INFO["note"])
